"""Create the Thai report in editable DOCX and Markdown from committed results.

Use the bundled Codex Python runtime (python-docx + Pillow) for this authoring
script. Export the reviewed DOCX with Word or the packaged document renderer.
The script does not read .env, create API clients, or make network requests.
"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parents[1]
FONT='Leelawadee UI'


def architecture():
    image=Image.new('RGB',(1440,580),'white')
    draw=ImageDraw.Draw(image)
    font_path=Path('C:/Windows/Fonts/arial.ttf')
    font=ImageFont.truetype(str(font_path),25) if font_path.exists() else ImageFont.load_default()
    small=ImageFont.truetype(str(font_path),21) if font_path.exists() else font
    def box(x,y,w,h,title,note=''):
        draw.rounded_rectangle((x,y,x+w,y+h),radius=16,fill='#f0f3ed',outline='#8e9b83',width=2)
        draw.text((x+w/2,y+25),title,font=font,anchor='mt',fill='#20271c')
        if note:draw.text((x+w/2,y+63),note,font=small,anchor='mt',fill='#5c6654')
    def arrow(a,b):
        draw.line((a,b),fill='#536347',width=3)
        import math
        angle=math.atan2(b[1]-a[1],b[0]-a[0])
        pts=[b,(b[0]-14*math.cos(angle-.45),b[1]-14*math.sin(angle-.45)),(b[0]-14*math.cos(angle+.45),b[1]-14*math.sin(angle+.45))]
        draw.polygon(pts,fill='#536347')
    box(20,30,300,110,'Thai UI + FastAPI','validated input')
    box(400,30,580,110,'gemma-4-E4B-it / class proxy','OpenAI SDK / native tool loop')
    box(1060,30,360,110,'Final JSON','schema + source checks')
    arrow((320,85),(400,85));arrow((980,85),(1060,85))
    box(450,215,480,100,'Tool dispatcher','allow-list / argument validation')
    arrow((690,140),(690,215))
    box(120,405,520,110,'search_docs','lexical TF-IDF / policy records')
    box(800,405,520,110,'query_inventory','SQLite SELECT / bound parameters')
    arrow((570,315),(380,405));arrow((810,315),(1060,405))
    draw.text((720,550),'Tool results return to the model as role: tool with tool_call_id',font=small,anchor='mm',fill='#5c6654')
    path=ROOT/'docs/architecture.png';image.save(path)
    return path


def font_run(run,size=None,bold=None):
    run.font.name=FONT
    rpr=run._element.get_or_add_rPr()
    fonts=rpr.rFonts
    if fonts is None:
        fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
    for script in ('ascii','hAnsi','eastAsia','cs'):
        fonts.set(qn('w:'+script),FONT)
    if size:run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    run.font.color.rgb=RGBColor(0,0,0)


def set_style(style,size,bold=False):
    style.font.name=FONT;style.font.size=Pt(size);style.font.bold=bold
    style.font.color.rgb=RGBColor(0,0,0)
    fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
    for script in ('ascii','hAnsi','eastAsia','cs'):fonts.set(qn('w:'+script),FONT)
    style.paragraph_format.line_spacing=1.2
    style.paragraph_format.space_after=Pt(6)


def add_table(doc,rows,widths):
    table=doc.add_table(rows=1,cols=len(rows[0]))
    table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    for i,w in enumerate(widths):table.columns[i].width=Cm(w)
    for ri,values in enumerate(rows):
        cells=table.rows[0].cells if ri==0 else table.add_row().cells
        for i,value in enumerate(values):
            cells[i].width=Cm(widths[i]);cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cellpr=cells[i]._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                item=OxmlElement('w:'+edge);item.set(qn('w:val'),'single');item.set(qn('w:sz'),'4');item.set(qn('w:color'),'D9D9D9');borders.append(item)
            cellpr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                item=OxmlElement('w:'+edge);item.set(qn('w:w'),'90');item.set(qn('w:type'),'dxa');margins.append(item)
            cellpr.append(margins)
            if ri==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E9EEF3');cellpr.append(shading)
            p=cells[i].paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.12
            if i>0 and len(str(value))<18:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            font_run(p.add_run(str(value)),9.3,ri==0)
        trpr=table.rows[ri]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)


def build_content():
    results=json.loads((ROOT/'results/retrieval.json').read_text(encoding='utf-8'))
    rows=[['วิธีค้นข้อมูล','k','Recall','Precision','MRR','ครบทุกแหล่ง']]
    names={'vector_top1':'Vector top 1 baseline','vector_top3':'Vector top 3 เลือกใช้','bm25_top3':'BM25 top 3','hybrid_top3':'Hybrid RRF top 3'}
    for s in results['systems']:
        rows.append([names[s['system']],str(s['cutoff']),f"{s['recall']:.4f}",f"{s['precision']:.4f}",f"{s['mrr']:.4f}",f"{s['complete']:.4f}"])
    return [
      {'title':'รายงานโครงงาน GearMate','blocks':[
        ('subtitle','ผู้ช่วยค้นกฎการยืมและตรวจอุปกรณ์ชมรม'),
        ('p','ผู้พัฒนา Chayapon8251\nRepository https://github.com/Chayapon8251/llm_lastwork\nวันที่ทดสอบ 25 ตุลาคม 2569'),
        ('h','1 ปัญหาและผู้ใช้งาน'),
        ('p','ผมเลือกทำ GearMate สำหรับนักศึกษาที่ต้องการยืมอุปกรณ์ชมรม เพราะคำถามหนึ่งมักต้องใช้ข้อมูลมากกว่าหนึ่งแบบ เช่น “กล้องว่างกี่ตัว และต้องวางเงินประกันเท่าไร” ส่วนแรกเป็นสถานะในตาราง ส่วนหลังเป็นเงื่อนไขในคู่มือ ถ้าให้โมเดลตอบจากความจำอย่างเดียวก็อาจได้จำนวนหรือกฎที่ไม่ตรงข้อมูล'),
        ('p','ระบบจึงให้ผู้ใช้ถามผ่านหน้าเว็บภาษาไทย แล้วให้ agent เลือกค้นคู่มือด้วย search_docs หรืออ่านรายการอุปกรณ์ด้วย query_inventory ผลลัพธ์แสดงคำตอบ แหล่งอ้างอิง และขั้นตอนการเรียก tool เพื่อให้ตรวจสอบได้ว่าใช้ข้อมูลอะไร'),
        ('h','2 ขอบเขตข้อมูลและสิ่งที่ทำได้'),
        ('p','ข้อมูลเป็นชมรมสมมติที่สร้างขึ้นสำหรับโครงงานนี้ ประกอบด้วยนโยบาย 18 หัวข้อและอุปกรณ์ 12 ชิ้น ไม่ใช่นโยบายหรือสถานะจริงของมหาวิทยาลัย และไม่ได้ใช้ corpus ของแอปตัวอย่างรายวิชา ระบบอ่านข้อมูลอย่างเดียว ไม่มีการจอง ชำระเงิน หรือแก้ไขรายการยืม'),
        ('p','วันที่ทดสอบสามารถเรียก class proxy ด้วยโมเดล gemma-4-E4B-it ได้จริงแล้ว โดยลองคำทักทาย คำถามที่ต้องใช้ทั้ง SQL และคู่มือ และคำถามเรื่องโดรนที่ระบบไม่รองรับ ส่วนการประเมิน retrieval ใช้ชุดคำถามติด label 30 ข้อแยกต่างหาก'),
        ('h','3 ข้อสรุปของการออกแบบ'),
        ('p','ผมเลือก vector top-3 เป็นค่าเริ่มต้น เพราะ recall เพิ่มจาก baseline top-1 ส่วน hybrid ให้ผลเท่ากันเมื่อเปรียบเทียบที่ k=3 จึงยังไม่มีเหตุผลจากชุดทดลองนี้ให้เพิ่มความซับซ้อนเป็นค่าเริ่มต้น ตัวค้นปัจจุบันใช้ TF-IDF แบบ lexical ไม่ใช่ neural semantic embeddings')
      ]},
      {'title':'สถาปัตยกรรมและเทคนิคที่ใช้','blocks':[
        ('h','4 โครงสร้างระบบ'),('figure','architecture'),
        ('p','ภาพที่ 1 หน้าเว็บส่งคำถามเข้า FastAPI จากนั้น agent ขอใช้ tool ผ่าน class proxy โค้ดตรวจชื่อและ arguments ก่อนรัน แล้วส่งผลกลับให้โมเดล คำตอบขั้นสุดท้ายถูกตรวจรูปแบบและแหล่งอ้างอิงอีกชั้น'),
        ('h','5 วงรอบของ agent'),
        ('p','เริ่มจากรับคำถามและบริบทที่จำกัดความยาว ส่งให้ OpenAI() พร้อม tool schemas โดยใช้ model="gemma-4-E4B-it" และ temperature=0 หากโมเดลส่ง tool_calls โค้ดจะเก็บ assistant message แปลง arguments จาก JSON ตรวจด้วย Pydantic แล้วรัน tool ที่อยู่ใน allow-list'),
        ('p','ผลแต่ละ call ถูกเพิ่มเป็น role: "tool" พร้อม tool_call_id ที่ตรงกัน รวมถึงเมื่อมีหลาย call ในรอบเดียว โมเดลจึงอ่านผลแล้วเลือกว่าจะค้นต่อหรือสรุปคำตอบ หยุดตามปกติเมื่อไม่มี tool_calls และมี max_steps=5 กับงบรวม 8 tool calls เป็นตัวหยุดสำรอง รวมทั้งปฏิเสธการเรียกเดิมซ้ำ'),
        ('h','6 ความเชื่อมโยงกับเนื้อหารายวิชา'),
        ('p','โครงงานใช้ structured output พร้อม schema validation, evaluation บนชุดติด label, agent loop ที่มี function calling และ stop condition, SQL แบบอ่านอย่างเดียว และ FastAPI REST/UI จึงมีทั้ง evaluation และความสามารถแบบ agentic ตามขอบเขตงาน'),
        ('p','ขั้นเลือก tool ใช้ native tool calling ส่วนขั้นจัดคำตอบเป็น JSON ใช้การเรียกอีกครั้งที่ไม่มี tools เพื่อไม่ให้คำสั่ง “ตอบเฉพาะ JSON” รบกวนการเลือก tool คำตอบถูกตรวจให้มี answer, source_ids และ supported และ source_ids ต้องมาจาก tool ที่ถูกเรียกจริง')
      ]},
      {'title':'ข้อมูล การค้น และขอบเขตความปลอดภัย','blocks':[
        ('h','7 การเตรียม corpus และ index'),
        ('p','แต่ละหัวข้อนโยบายเป็นหนึ่ง chunk เพื่อไม่ตัดเงื่อนไขกลางเรื่อง เช่น การต่อเวลาต้องแจ้งก่อน 24 ชั่วโมงและได้รับอนุมัติ ข้อมูลมี ID คงที่ ข้อความไทย และคำแปลอังกฤษ เพื่อใช้อ้างอิงและทดสอบทั้งสองภาษา ไฟล์ดิบอยู่ใน data/policies.json ส่วน snapshot อุปกรณ์อยู่ใน data/inventory.json'),
        ('p','ตัว vector ใช้ character trigrams แปลงเป็น TF-IDF แล้วคำนวณ cosine similarity จึงไม่ต้องใช้บริการ embedding ภายนอก ส่วน BM25 ใช้ English tokens ร่วมกับ Thai bigrams ทั้งสอง index สร้างในหน่วยความจำจากข้อมูลใน repo จุดอ่อนคือยังอาศัยคำหรือตัวอักษรที่ตรงกัน และอาจพลาดคำถามที่ใช้คำต่างจากเอกสารมาก'),
        ('h','8 วิธีที่นำมาเปรียบเทียบ'),
        ('p','ทดลอง vector top-1 เป็น baseline แบบง่าย จากนั้น vector top-3 เพื่อดูผลของจำนวน context, BM25 top-3 และ hybrid top-3 ที่รวมอันดับด้วย RRF โดย score(doc)=sum(1/(60+rank)) เริ่ม rank ที่ 1 และนำ top-10 ของแต่ละตัวค้นมารวมกัน ไม่เฉลี่ยคะแนน BM25 กับ cosine เพราะอยู่คนละสเกล'),
        ('p','ใช้ cosine floor 0.04 เหมือนกันทุกวิธี ค่านี้เป็น heuristic ระหว่างพัฒนา ไม่ใช่คะแนนความมั่นใจที่ปรับเทียบแล้ว ยังไม่มีการทดสอบ negative set ขนาดใหญ่ และไม่ควรสรุปว่าผ่าน threshold แล้วคำตอบจะถูกเสมอ เวอร์ชันนี้ยังไม่มี cross-encoder หรือ LLM reranker'),
        ('h','9 การควบคุม SQL และข้อมูลจาก tool'),
        ('p','query_inventory รับเฉพาะ category, code และ status โค้ดเป็นผู้สร้าง SELECT จาก template ใช้ ? bind values และเปิด SQLite เป็น query_only พร้อม LIMIT 30 จึงไม่มีช่องให้โมเดลส่ง SQL อิสระหรือเปลี่ยนตาราง ตัว validator ปฏิเสธชื่อหมวดและ arguments นอก schema'),
        ('p','ข้อความจาก history และผล tool ถือเป็นข้อมูลที่ไม่เชื่อถือ ไม่ใช่คำสั่งให้ข้ามข้อจำกัด ระบบไม่มี tool สำหรับจอง ตัดเงิน ลบข้อมูล หรือรันโค้ดทั่วไป ส่วนการตรวจ citation ยืนยันได้เพียงว่า ID เคยถูกค้นจริง ยังต้องตรวจว่าข้อความในคำตอบสอดคล้องกับหลักฐานด้วย')
      ]},
      {'title':'การประเมิน retrieval และผลเปรียบเทียบ','blocks':[
        ('h','10 ชุดข้อมูลและวิธีวัด'),
        ('p','ชุด eval/test.json มี 30 คำถาม แบ่งเป็นคำถามที่ต้องใช้แหล่งเดียว 18 ข้อและสองแหล่ง 12 ข้อ มีคำถามไทย รหัสอุปกรณ์ตรงตัว และคำถามอังกฤษหลายประเด็น มีชุด dev อีก 6 ข้อและสถานการณ์ที่ข้อมูลไม่รองรับอีก 6 ข้อแยกไว้ ไม่รวมกรณีไม่มี relevant เข้าไปทำให้ค่า recall ดูดีขึ้น'),
        ('p','labels สร้างควบคู่กับ corpus โดยใช้ AI ช่วย ยังไม่ได้ให้ผู้ประเมินอิสระตรวจ จึงมีความเสี่ยงที่คำในคำถามตรงกับเอกสารมากเกินไป และหลังจากเปิดดูผลแล้วควรถือชุดนี้เป็น development benchmark ก่อนอ้างผลทั่วไปต้องสร้าง held-out set ใหม่'),
        ('p','Recall@k วัดสัดส่วน relevant IDs ที่อยู่ในผล k อันดับแรก Precision@k หารด้วย k ส่วน MRR เฉลี่ย 1/rank ของ relevant ตัวแรกในรายการที่คืนมา ถ้าไม่พบให้เป็น 0 และวัดสัดส่วนคำถามที่ค้นครบทุกแหล่งเพิ่มด้วย ทุก metric เฉลี่ยรายคำถามและระบุ k ของแต่ละระบบ'),
        ('table',rows),
        ('p','ตารางที่ 1 ผลที่รันซ้ำได้จาก python -m scripts.evaluate โดยไม่มีการเรียก LLM'),
        ('h','11 การตีความและการเลือกค่าเริ่มต้น'),
        ('p','จาก top-1 เป็น top-3 recall เพิ่มจาก 0.8000 เป็น 0.9833 หรือประมาณ 18.33 percentage points และค้นครบทุกแหล่งเพิ่มจาก 60.00% เป็น 96.67% แต่ precision ลดเพราะมีบริบทที่ไม่เกี่ยวข้องเพิ่มขึ้น ความดีขึ้นนี้ส่วนหนึ่งมาจากงบจำนวน chunk ไม่ใช่หลักฐานว่าอัลกอริทึม hybrid ดีกว่า'),
        ('p','เมื่อเทียบที่ k=3 เท่ากัน hybrid เสมอกับ vector จึงเลือก vector top-3 ส่วน MRR=1.0 แปลว่า relevant แรกอยู่ลำดับแรกในชุดนี้ ไม่ได้หมายความว่าหลักฐานครบหรือคำตอบ LLM ถูกทุกข้อ ไฟล์ results/retrieval.json เก็บอันดับรายข้อ failure และ hash ของข้อมูลกับโค้ดเพื่อทำซ้ำ')
      ]},
      {'title':'ผลเรียกโมเดลจริงและกรณีที่ยังผิดพลาด','blocks':[
        ('h','12 การทดสอบ class proxy'),
        ('p','หลังตั้ง environment ใน .env ได้ทดสอบผ่าน OpenAI SDK กับ class proxy จริงโดยใช้ gemma-4-E4B-it สามสถานการณ์ต่อไปนี้ เป็น smoke test เพื่อตรวจการเชื่อมต่อและเส้นทาง tool ไม่ใช่การประเมิน faithfulness หรือ accuracy ที่ครอบคลุมผู้ใช้จริง'),
        ('table',[
          ['คำถาม','สิ่งที่เกิดขึ้น','เวลารวมโดยประมาณ'],
          ['สวัสดีครับ','ตอบทักทายโดยไม่เรียก tool','2.93 วินาที'],
          ['กล้องว่างกี่ตัว และเงินประกันเท่าไร','ใช้ query_inventory และ search_docs ตอบ 2 กล้อง ประกัน 300 บาทต่อรายการ อ้าง inventory และ camera_deposit','2.04 วินาที'],
          ['มีโดรนให้ยืมไหม','ขอ query_inventory ด้วย category=drone ถูก validator ปฏิเสธ แล้วตอบว่าไม่มีข้อมูลในหมวดที่รองรับ','1.17 วินาที']
        ]),
        ('p','เวลานี้เป็นผลรอบเดียวต่อสถานการณ์ รวมการจัดคำตอบเป็น JSON ไม่ใช่ค่าเฉลี่ย latency และอาจเปลี่ยนเมื่อรันใหม่ ผลโดรนแสดงว่าระบบป้องกัน arguments นอก schema ได้ แต่ไม่ได้พิสูจน์ว่ามีการค้นฐานข้อมูลแล้วพบโดรนเป็นศูนย์ จึงควรใช้ถ้อยคำว่า “ระบบไม่รองรับหมวดนี้”'),
        ('h','13 ข้อผิดพลาดในการค้นที่ต้องรายงาน'),
        ('p','ข้อ t27 ถามทั้งเงินประกันกล้องและการอบรม ต้องการ camera_deposit กับ training แต่ vector top-3 คืน training, lighting_deposit และ audio_deposit จึงได้ recall 0.5 แม้ MRR ยังเป็น 1 เพราะ relevant ตัวแรกอยู่บนสุด คำเรื่องเงินประกันหลายประเภททำให้ตัวค้นสับสน'),
        ('p','อีกกรณีที่ตรวจหลังพัฒนาคือคำถามสั้น “ยืมได้กี่วัน” ตัวค้นตรง ๆ ได้ extension, membership และ weekend_return แทน duration แต่คำว่า “ระยะเวลายืมปกติ” ค้นเจอ กรณีนี้เป็นข้อจำกัดของ retriever ไม่ใช่ผลทดลองว่า live agent ล้มเหลวเสมอ เพราะโมเดลอาจปรับ query ก่อนค้น'),
        ('p','การทดสอบอัตโนมัติ 39 tests ครอบคลุม validators, SQL query-only, bound parameters, API และ protocol ของ agent ด้วย stubbed responses ต้องแยกผลนี้ออกจาก smoke test ที่เรียกโมเดลจริง และยังต้องมีคนตรวจว่าคำตอบอ้างหลักฐานได้ตรงหรือไม่')
      ]},
      {'title':'ข้อจำกัด สิ่งที่ทำต่อ และแหล่งอ้างอิง','blocks':[
        ('h','14 ข้อจำกัดของผลทดลอง'),
        ('p','ข้อมูลและ snapshot มีขนาดเล็กและเป็นข้อมูลจำลอง ไม่มีการเชื่อมระบบชมรมจริง ชุดคำถามสร้างจาก corpus เดียวกันและยังไม่มี independent labels จึงไม่ควรใช้ recall สูงในชุดนี้ไปอ้างว่าตอบคำถามจริงได้แม่นยำเท่ากัน ตัว TF-IDF เป็น lexical representation และ cosine floor ยังไม่ผ่านการปรับเทียบกับกรณีนอกโดเมนจำนวนมาก'),
        ('p','ผลโมเดลจริงที่บันทึกมีเพียงสามสถานการณ์ แสดงว่าเชื่อม proxy และเรียก tool ได้ ไม่พอจะสรุป accuracy, robustness หรือ faithfulness โดยรวม การมี citation อยู่ในผล tool ก็ไม่ได้รับประกันว่าทุกประโยคของคำตอบถูกต้อง'),
        ('h','15 สิ่งที่ควรทำต่อก่อนส่งงาน'),
        ('p','ผมควรตรวจ corpus กับ labels เองและให้คนอีกคนเขียนคำถามใหม่ เพิ่มการตรวจคำตอบ LLM เป็นรายข้อว่าถูกครบ มีข้ออ้างที่ไม่มีหลักฐาน หรือควรปฏิเสธ จากนั้นทดลอง multilingual embeddings, query refinement หรือ reranker โดยคุมจำนวน chunk เท่ากัน และเก็บเฉพาะการเปลี่ยนแปลงที่คุ้มกับ latency และค่าใช้จ่าย'),
        ('p','ก่อน demo ต้องอ่านและอธิบายเส้นทาง tool, เหตุผลที่เลือก top-3 และความต่างระหว่าง recall กับ MRR ได้ พร้อมลองคำถามใหม่ที่ไม่ได้เตรียมไว้ ส่วนคลิปสาธิต 3-6 นาทียังไม่ได้บันทึก และจะแสดงทั้งกรณีสำเร็จกับข้อจำกัดตามโจทย์ในขั้นต่อไป'),
        ('h','16 การใช้ AI และสิ่งที่ผู้ส่งต้องรับผิดชอบ'),
        ('p','โค้ด corpus และร่างเอกสารนี้พัฒนาด้วยความช่วยเหลือจาก AI ตามที่โจทย์อนุญาต การแก้สำนวนควรทำเพื่อให้ตรงความเข้าใจของผู้ส่ง ไม่ควรเปลี่ยนผลทดลองหรือเขียนอ้างว่าทดสอบกับผู้ใช้จริงทั้งที่ยังไม่ได้ทำ ผู้ส่งยังต้องตรวจความถูกต้องและอธิบายการออกแบบด้วยตนเอง'),
        ('h','17 แหล่งอ้างอิงและไฟล์สำหรับทำซ้ำ'),
        ('p','แนวคิด metrics และ RRF มาจาก Lecture 8 ส่วน tool calling, bounded agent loop และ SQL guardrails มาจาก Lectures 10-14 การต่อ native calls อ้างอิง OpenAI Function calling documentation ที่ https://developers.openai.com/api/docs/guides/function-calling'),
        ('p','Repository https://github.com/Chayapon8251/llm_lastwork\nรัน retrieval evaluation ด้วย python -m scripts.evaluate\nรัน tests ด้วย python -m pytest -q\nผล retrieval อยู่ที่ results/retrieval.json และผล smoke test ที่ตัด credentials ออกอยู่ที่ results/live-smoke.json')
      ]}
    ]


def main():
    sections=build_content()
    architecture_path=architecture()
    doc=Document()
    section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(1.9);section.bottom_margin=Cm(1.9)
    section.left_margin=Cm(2);section.right_margin=Cm(2)
    set_style(doc.styles['Normal'],11)
    set_style(doc.styles['Title'],23,True)
    set_style(doc.styles['Subtitle'],13)
    set_style(doc.styles['Heading 1'],16,True)
    set_style(doc.styles['Heading 2'],12.5,True)
    for name in ('Heading 1','Heading 2'):
        doc.styles[name].paragraph_format.space_before=Pt(7)
        doc.styles[name].paragraph_format.space_after=Pt(5)
    normal=doc.styles['Normal']
    normal.paragraph_format.widow_control=True
    # Thai language enables the renderer's normal Thai word-break behaviour.
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'th-TH');lang.set(qn('w:eastAsia'),'th-TH');normal.element.get_or_add_rPr().append(lang)
    doc.core_properties.title='รายงานโครงงาน GearMate'
    doc.core_properties.author='Chayapon8251'
    doc.core_properties.subject='ผู้ช่วยค้นกฎการยืมและตรวจอุปกรณ์ชมรม'
    markdown=[]
    for i,part in enumerate(sections):
        if i:doc.add_page_break()
        doc.add_paragraph(part['title'],style='Title' if i==0 else 'Heading 1')
        markdown.append(('# ' if i==0 else '## ')+part['title'])
        for kind,value in part['blocks']:
            if kind=='h':
                doc.add_paragraph(value,style='Heading 2');markdown.append('### '+value)
            elif kind=='subtitle':
                doc.add_paragraph(value,style='Subtitle');markdown.append(value)
            elif kind=='p':
                doc.add_paragraph(value);markdown.append(value)
            elif kind=='figure':
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(architecture_path),width=Cm(16.7))
                markdown.append('![สถาปัตยกรรมระบบ](architecture.png)')
            elif kind=='table':
                widths=[4.7,1,2.4,2.4,2.2,3.1] if len(value[0])==6 else [4.6,9.2,2]
                add_table(doc,value,widths)
                markdown.append('\n'.join(['| '+' | '.join(row)+' |' for row in value[:1]]+['| '+' | '.join(['---']*len(value[0]))+' |']+['| '+' | '.join(row)+' |' for row in value[1:]]))
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    font_run(footer.add_run('GearMate   หน้า '),8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.save(ROOT/'docs/report-th.docx')
    (ROOT/'docs/report-th.md').write_text('\n\n'.join(markdown)+'\n',encoding='utf-8')
    print('Created docs/report-th.docx and docs/report-th.md')


if __name__=='__main__':
    main()
